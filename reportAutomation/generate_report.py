"""
Production vs Captured Count Report Generator
Compares DB production counts with captured counts from model logs,
sensor trigger data, and image folder counts.

Usage (CLI):
  py generate_report.py                           (interactive date input)
  py generate_report.py 14-03-2026 18-03-2026     (date range DD-MM-YYYY)
  py generate_report.py "14-03-2026 06:00" "18-03-2026 22:00"  (with time)
  py generate_report.py 14-03-2026 18-03-2026 --csv path/to/file.csv  (use CSV instead of DB)

Usage (Debugger / direct function call):
  generate_report("14-03-2026", "18-03-2026")
  generate_report("14-03-2026 06:00", "18-03-2026 22:00")
  generate_report("14-03-2026", "18-03-2026", csv_path="path/to/file.csv")

By default, fetches production data directly from SQL Server.
Use --csv (CLI) or csv_path parameter (function) to fall back to a CSV file.

All dates/times are in IST.
Outputs an Excel (.xlsx) report grouped by Date and Shift.
"""

import json
import re
import os
import sys
from datetime import datetime, timedelta
from collections import defaultdict

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl not installed. Installing...")
    os.system(f"{sys.executable} -m pip install openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

# -- Configuration ----------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

SINGLE_PLATE_LOG = os.path.join(BASE_DIR, "plate", "SinglePlateModel.log")
MULTI_PLATE_LOG = os.path.join(BASE_DIR, "bunch", "MultiplatePlateModel.log")
SENSOR_PLATE_JSON = os.path.join(BASE_DIR, "plate", "sensorTrigerPlate.json.bak")
SENSOR_BUNCH_JSON = os.path.join(BASE_DIR, "bunch", "sensorTrigerBunch.json.bak")
DOUBLE_PLATE_DIR = os.path.join(BASE_DIR, "plate", "doublePlateImages")
NEG_IMAGES_DIR = os.path.join(BASE_DIR, "plate", "NegImages")
OUTPUT_DIR = os.path.join(BASE_DIR, "reportAutomation")

# -- SQL Server Connection --------------------------------------------------
DB_CONNECTION_STRING = (
    "DRIVER={ODBC Driver 17 for SQL Server};"
    "SERVER=<YOUR_SERVER>;"
    "DATABASE=areml;"
    "UID=<YOUR_UID>;"
    "PWD=<YOUR_PWD>"
)
DB_TABLE = "[areml].[dbo].[ABD2_BW2]"

# IST = UTC + 5:30
IST_OFFSET = timedelta(hours=5, minutes=30)

SHIFTS = {
    "A": (6, 14),   # 06:00 to 14:00
    "B": (14, 22),  # 14:00 to 22:00
    "C": (22, 6),   # 22:00 to 06:00 (crosses midnight)
}


# -- Date range parsing -----------------------------------------------------

def parse_date_input(text):
    """Parse a date/datetime string in IST. Accepts DD-MM-YYYY or DD-MM-YYYY HH:MM."""
    text = text.strip()
    for fmt in ("%d-%m-%Y %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise ValueError(f"Cannot parse date: '{text}'. Use DD-MM-YYYY or DD-MM-YYYY HH:MM")


def get_date_range_and_csv_from_args():
    """
    Get IST date range from CLI args or interactive input.
    Also checks for --csv flag to use CSV file instead of DB.
    Returns (from_dt, to_dt, csv_path_or_None).
    """
    csv_path = None
    args = sys.argv[1:]

    # Extract --csv flag if present
    if "--csv" in args:
        csv_idx = args.index("--csv")
        if csv_idx + 1 < len(args):
            csv_path = args[csv_idx + 1]
        else:
            print("Error: --csv requires a file path argument")
            sys.exit(1)
        args = args[:csv_idx] + args[csv_idx + 2:]

    if len(args) >= 2:
        from_dt = parse_date_input(args[0])
        to_dt = parse_date_input(args[1])
    else:
        print("\nEnter date range in IST (format: DD-MM-YYYY or DD-MM-YYYY HH:MM)")
        from_str = input("  From: ").strip()
        to_str = input("  To  : ").strip()
        from_dt = parse_date_input(from_str)
        to_dt = parse_date_input(to_str)

    # If only date given (no time), default from=00:00 to=23:59:59
    if from_dt.hour == 0 and from_dt.minute == 0 and from_dt.second == 0:
        pass  # already at start of day
    if to_dt.hour == 0 and to_dt.minute == 0 and to_dt.second == 0:
        to_dt = to_dt.replace(hour=23, minute=59, second=59)

    return from_dt, to_dt, csv_path


def is_in_range(ist_dt, from_dt, to_dt):
    """Check if an IST datetime falls within the range [from_dt, to_dt]."""
    return from_dt <= ist_dt <= to_dt


# -- Helpers ----------------------------------------------------------------

def utc_to_ist(dt):
    return dt + IST_OFFSET


def get_shift_and_date(ist_dt):
    """
    Given an IST datetime, return (shift_letter, shift_date).
    For shift C (22:00-06:00), if time is 00:00-06:00, shift_date is previous day.
    """
    hour = ist_dt.hour
    if 6 <= hour < 14:
        return "A", ist_dt.date()
    elif 14 <= hour < 22:
        return "B", ist_dt.date()
    else:
        if hour >= 22:
            return "C", ist_dt.date()
        else:
            return "C", (ist_dt - timedelta(days=1)).date()


# -- Parsers ----------------------------------------------------------------

def parse_model_log(filepath, from_dt, to_dt):
    """
    Parse SinglePlateModel.log or MultiplatePlateModel.log.
    Timestamps are in UTC -> converted to IST, then filtered by range.
    Returns list of (ist_datetime, moved, deleted, total_captured).
    """
    results = []
    if not os.path.exists(filepath):
        print(f"  Warning: {filepath} not found")
        return results

    batch_pattern = re.compile(
        r'\[(\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})\].*?'
        r'Moved:\s*(\d+).*?Deleted:\s*(\d+)'
    )

    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        for line in f:
            m = batch_pattern.search(line)
            if m:
                utc_dt = datetime.strptime(m.group(1), "%Y-%m-%d %H:%M:%S")
                ist_dt = utc_to_ist(utc_dt)
                if not is_in_range(ist_dt, from_dt, to_dt):
                    continue
                moved = int(m.group(2))
                deleted = int(m.group(3))
                results.append((ist_dt, moved, deleted, moved + deleted))

    return results


def parse_sensor_trigger_json(filepath, from_dt, to_dt):
    """
    Parse sensor trigger JSON. Timestamps are in IST.
    Only includes entries that overlap with [from_dt, to_dt].
    For partial overlaps, clips the entry to the range.
    Returns list of (start_ist, end_ist, trigger_count, duration_seconds).
    """
    results = []
    if not os.path.exists(filepath):
        print(f"  Warning: {filepath} not found")
        return results

    with open(filepath, 'r', encoding='utf-8') as f:
        data = json.load(f)

    for key, entry in data.items():
        start_dt = datetime.strptime(entry["start_dateTime"], "%Y-%m-%dT%H:%M:%S")
        end_dt = datetime.strptime(entry["end_dateTime"], "%Y-%m-%dT%H:%M:%S")
        trigger_count = entry["trigger_count"]
        duration = entry["duration_seconds"]

        # Skip if entirely outside range
        if end_dt < from_dt or start_dt > to_dt:
            continue

        # Clip to range and adjust trigger count proportionally
        orig_start, orig_end = start_dt, end_dt
        clipped_start = max(start_dt, from_dt)
        clipped_end = min(end_dt, to_dt)

        total_secs = (orig_end - orig_start).total_seconds()
        if total_secs > 0:
            clipped_secs = (clipped_end - clipped_start).total_seconds()
            proportion = clipped_secs / total_secs
            adj_count = round(trigger_count * proportion)
        else:
            adj_count = trigger_count

        results.append((clipped_start, clipped_end, adj_count, duration))

    return results


def parse_db_csv(filepath, from_dt, to_dt):
    """
    Parse the DB CSV. Datetime is IST (DD-MM-YYYY HH:MM). Filtered by range.
    Returns list of (ist_datetime, shift_letter, bunch_qty).
    """
    results = []
    if not os.path.exists(filepath):
        print(f"  Warning: {filepath} not found")
        return results

    with open(filepath, 'r', encoding='utf-8') as f:
        header = f.readline().strip().split(',')
        dt_idx = header.index("Datetime")
        shift_idx = header.index("Shift")
        qty_idx = header.index("Bunch_Qty")

        for line in f:
            parts = line.strip().split(',')
            if len(parts) <= max(dt_idx, shift_idx, qty_idx):
                continue
            try:
                dt_str = parts[dt_idx].strip()
                ist_dt = datetime.strptime(dt_str, "%d-%m-%Y %H:%M")
                if not is_in_range(ist_dt, from_dt, to_dt):
                    continue
                shift = parts[shift_idx].strip()
                bunch_qty = int(parts[qty_idx].strip())
                results.append((ist_dt, shift, bunch_qty))
            except (ValueError, IndexError):
                continue

    return results


def fetch_db_data(from_dt, to_dt):
    """
    Fetch production data directly from SQL Server.
    The DB Datetime column is in IST. Filtered by the given IST range.
    Returns list of (ist_datetime, shift_letter, bunch_qty) — same format as parse_db_csv.
    """
    try:
        import pyodbc
    except ImportError:
        print("  pyodbc not installed. Installing...")
        os.system(f"{sys.executable} -m pip install pyodbc")
        import pyodbc

    # Format dates for SQL Server query (YYYY-MM-DD HH:MM:SS)
    from_str = from_dt.strftime("%Y-%m-%d %H:%M:%S")
    to_str = to_dt.strftime("%Y-%m-%d %H:%M:%S")

    query = f"""
        SELECT Datetime, Shift, Bunch_Qty
        FROM {DB_TABLE}
        WHERE Datetime >= ?
        AND Datetime <= ?
        ORDER BY Datetime ASC
    """

    results = []
    conn = None
    try:
        conn = pyodbc.connect(DB_CONNECTION_STRING, timeout=10)
        cursor = conn.cursor()
        cursor.execute(query, from_str, to_str)

        for row in cursor.fetchall():
            ist_dt = row[0]  # pyodbc returns datetime objects directly
            if isinstance(ist_dt, str):
                ist_dt = datetime.strptime(ist_dt, "%Y-%m-%d %H:%M:%S")
            shift = str(row[1]).strip()
            bunch_qty = int(row[2])
            results.append((ist_dt, shift, bunch_qty))

        cursor.close()
    except pyodbc.Error as e:
        print(f"  DB Error: {e}")
        print("  Tip: Check DB_CONNECTION_STRING in generate_report.py")
        print("  Or use --csv flag to fall back to CSV file.")
        sys.exit(1)
    finally:
        if conn:
            conn.close()

    return results


def count_images_in_range(directory, from_dt, to_dt):
    """
    Count images recursively in directory and all subfolders.
    Only counts images whose filename timestamp falls within [from_dt, to_dt].
    Image timestamps (in filename) are in IST.
    Returns (shift_counts_dict, total_in_range, total_all).
    """
    counts = defaultdict(int)
    total_in_range = 0
    total_all = 0
    if not os.path.exists(directory):
        return counts, total_in_range, total_all

    ts_pattern = re.compile(r'_(\d{14})')  # YYYYMMDDHHMMSS

    for root, dirs, files in os.walk(directory):
        for fname in files:
            if fname.lower().endswith(('.jpg', '.jpeg', '.png', '.bmp')):
                total_all += 1
                m = ts_pattern.search(fname)
                if m:
                    try:
                        ist_dt = datetime.strptime(m.group(1), "%Y%m%d%H%M%S")
                        if is_in_range(ist_dt, from_dt, to_dt):
                            total_in_range += 1
                            shift, shift_date = get_shift_and_date(ist_dt)
                            counts[(shift_date, shift)] += 1
                    except ValueError:
                        pass

    return counts, total_in_range, total_all


def distribute_sensor_triggers_by_shift(entries):
    """
    Distribute sensor trigger counts across shifts proportionally.
    Returns dict: (shift_date, shift) -> trigger_count
    """
    counts = defaultdict(float)

    for start_dt, end_dt, trigger_count, duration in entries:
        if start_dt == end_dt or (end_dt - start_dt).total_seconds() <= 0:
            shift, shift_date = get_shift_and_date(start_dt)
            counts[(shift_date, shift)] += trigger_count
            continue

        current = start_dt
        total_duration = (end_dt - start_dt).total_seconds()

        while current < end_dt:
            shift, shift_date = get_shift_and_date(current)

            if shift == "C":
                if current.hour >= 22:
                    shift_end = current.replace(
                        hour=0, minute=0, second=0, microsecond=0
                    ) + timedelta(days=1, hours=6)
                else:
                    shift_end = current.replace(
                        hour=6, minute=0, second=0, microsecond=0
                    )
            else:
                _, end_h = SHIFTS[shift]
                shift_end = current.replace(
                    hour=end_h, minute=0, second=0, microsecond=0
                )
                if shift_end <= current:
                    shift_end += timedelta(days=1)

            chunk_end = min(shift_end, end_dt)
            chunk_seconds = (chunk_end - current).total_seconds()
            proportion = chunk_seconds / total_duration
            counts[(shift_date, shift)] += trigger_count * proportion
            current = chunk_end

    return {k: round(v) for k, v in counts.items()}


# -- Excel Report -----------------------------------------------------------

def generate_report(from_date=None, to_date=None, csv_path=None):
    """
    Generate the production vs captured count report.

    Can be called in three ways:
      1. CLI:       py generate_report.py 14-03-2026 18-03-2026
      2. Function:  generate_report("14-03-2026", "18-03-2026")
      3. Function:  generate_report("14-03-2026 06:00", "18-03-2026 22:00", csv_path="file.csv")

    Parameters:
        from_date (str, optional): Start date in IST. Format: DD-MM-YYYY or DD-MM-YYYY HH:MM.
                                   If None, falls back to CLI args or interactive input.
        to_date   (str, optional): End date in IST. Same format as from_date.
        csv_path  (str, optional): Path to CSV file. If None, uses SQL Server.
    """
    print("=" * 60)
    print("  Production vs Captured Count Report Generator")
    print("=" * 60)

    # Get date range: prefer function params > CLI args > interactive
    if from_date is not None and to_date is not None:
        from_dt = parse_date_input(from_date)
        to_dt = parse_date_input(to_date)
        # If only date given (no time), default to full day
        if to_dt.hour == 0 and to_dt.minute == 0 and to_dt.second == 0:
            to_dt = to_dt.replace(hour=23, minute=59, second=59)
    else:
        from_dt, to_dt, cli_csv = get_date_range_and_csv_from_args()
        if csv_path is None:
            csv_path = cli_csv

    print(f"\n  Report range (IST): {from_dt.strftime('%d-%m-%Y %H:%M')} "
          f"to {to_dt.strftime('%d-%m-%Y %H:%M')}")

    # 1. Fetch DB data (SQL Server by default, CSV if --csv flag used)
    if csv_path:
        print(f"\n[1/6] Reading DB data from CSV: {csv_path}")
        db_data = parse_db_csv(csv_path, from_dt, to_dt)
    else:
        print("\n[1/6] Querying SQL Server database...")
        db_data = fetch_db_data(from_dt, to_dt)
    print(f"  Found {len(db_data)} bunch records in range")

    db_bunches = defaultdict(int)
    db_plates = defaultdict(int)
    for ist_dt, shift, bunch_qty in db_data:
        s, sd = get_shift_and_date(ist_dt)
        db_bunches[(sd, shift)] += 1
        db_plates[(sd, shift)] += bunch_qty

    # 2. Parse model logs
    print("\n[2/6] Parsing model logs (UTC -> IST)...")
    single_plate_data = parse_model_log(SINGLE_PLATE_LOG, from_dt, to_dt)
    multi_plate_data = parse_model_log(MULTI_PLATE_LOG, from_dt, to_dt)
    print(f"  SinglePlateModel: {len(single_plate_data)} batch records in range")
    print(f"  MultiplatePlateModel: {len(multi_plate_data)} batch records in range")

    captured_single = defaultdict(int)
    for ist_dt, moved, deleted, total in single_plate_data:
        shift, shift_date = get_shift_and_date(ist_dt)
        captured_single[(shift_date, shift)] += total

    captured_multi = defaultdict(int)
    for ist_dt, moved, deleted, total in multi_plate_data:
        shift, shift_date = get_shift_and_date(ist_dt)
        captured_multi[(shift_date, shift)] += total

    # 3. Parse sensor triggers
    print("\n[3/6] Parsing sensor trigger JSONs (IST)...")
    sensor_plate_data = parse_sensor_trigger_json(SENSOR_PLATE_JSON, from_dt, to_dt)
    sensor_bunch_data = parse_sensor_trigger_json(SENSOR_BUNCH_JSON, from_dt, to_dt)
    print(f"  Plate sensor: {len(sensor_plate_data)} trigger sessions in range")
    print(f"  Bunch sensor: {len(sensor_bunch_data)} trigger sessions in range")

    sensor_plate_by_shift = distribute_sensor_triggers_by_shift(sensor_plate_data)
    sensor_bunch_by_shift = distribute_sensor_triggers_by_shift(sensor_bunch_data)

    # 4. Count images (recursive, filtered by date range)
    print("\n[4/6] Counting images in doublePlateImages and NegImages...")
    dbl_counts, dbl_in_range, dbl_total = count_images_in_range(
        DOUBLE_PLATE_DIR, from_dt, to_dt
    )
    neg_counts, neg_in_range, neg_total = count_images_in_range(
        NEG_IMAGES_DIR, from_dt, to_dt
    )
    print(f"  Double plate images: {dbl_in_range} in range (of {dbl_total} total)")
    print(f"  Negative images: {neg_in_range} in range (of {neg_total} total)")

    # 5. Collect all shift keys
    print("\n[5/6] Building report data...")
    all_keys = set()
    for d in [db_bunches, db_plates, captured_single, captured_multi,
              sensor_plate_by_shift, sensor_bunch_by_shift,
              dbl_counts, neg_counts]:
        all_keys.update(d.keys())

    if not all_keys:
        print("  No data found in the specified range!")
        return

    shift_order = {"A": 0, "B": 1, "C": 2}
    sorted_keys = sorted(all_keys, key=lambda x: (x[0], shift_order.get(x[1], 9)))

    # 6. Generate Excel
    print("\n[6/6] Generating Excel report...")
    wb = Workbook()
    ws = wb.active
    ws.title = "Shift Report"

    # -- Styles --
    header_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    data_align = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    yellow_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')

    # Title
    ws.merge_cells('A1:O1')
    title_cell = ws['A1']
    title_cell.value = "Production vs Captured Count - Shift Report"
    title_cell.font = Font(name='Calibri', bold=True, size=14, color='2F5496')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # Date range subtitle
    ws.merge_cells('A2:O2')
    data_source = f"CSV: {csv_path}" if csv_path else "SQL Server (live)"
    ws['A2'].value = (
        f"Range: {from_dt.strftime('%d-%m-%Y %H:%M')} to "
        f"{to_dt.strftime('%d-%m-%Y %H:%M')} IST  |  "
        f"DB Source: {data_source}  |  "
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )
    ws['A2'].font = Font(name='Calibri', italic=True, size=10)
    ws['A2'].alignment = Alignment(horizontal='center')

    # Headers (row 4)
    headers = [
        "Date",                              # A (col 1)
        "Shift",                             # B (col 2)
        "Shift Time (IST)",                  # C (col 3)
        "DB Bunches",                        # D (col 4)
        "DB Plates\n(Bunches x Qty)",        # E (col 5)
        "Captured Plates\n(SinglePlate Log)", # F (col 6)
        "Captured Plates\n(MultiPlate Log)",  # G (col 7)
        "Sensor Triggers\n(Plate)",          # H (col 8)
        "Sensor Triggers\n(Bunch)",          # I (col 9)
        "Double Plate\nImages",              # J (col 10)
        "Negative\nImages",                  # K (col 11)
        "Plate Diff\n(DB - Captured)",       # L (col 12)
        "Plate\nStatus",                     # M (col 13)
        "Bunch Diff\n(DB - Captured)",       # N (col 14)
        "Bunch\nStatus",                     # O (col 15)
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[4].height = 40

    shift_times = {
        "A": "06:00 - 14:00",
        "B": "14:00 - 22:00",
        "C": "22:00 - 06:00",
    }

    # Data rows
    row = 5
    for shift_date, shift in sorted_keys:
        key = (shift_date, shift)
        db_b = db_bunches.get(key, 0)
        db_p = db_plates.get(key, 0)
        cap_s = captured_single.get(key, 0)
        cap_m = captured_multi.get(key, 0)
        sensor_p = sensor_plate_by_shift.get(key, 0)
        sensor_b = sensor_bunch_by_shift.get(key, 0)
        dbl_img = dbl_counts.get(key, 0)
        neg_img = neg_counts.get(key, 0)

        # Plate difference: DB Plates vs SinglePlate captured
        plate_diff = db_p - cap_s
        if plate_diff > 0:
            plate_status = "DB MORE"
        elif plate_diff < 0:
            plate_status = "CAPTURED MORE"
        else:
            plate_status = "NO DATA" if (db_p == 0 and cap_s == 0) else "MATCH"

        # Bunch difference: DB Bunches vs MultiPlate captured
        bunch_diff = db_b - cap_m
        if bunch_diff > 0:
            bunch_status = "DB MORE"
        elif bunch_diff < 0:
            bunch_status = "CAPTURED MORE"
        else:
            bunch_status = "NO DATA" if (db_b == 0 and cap_m == 0) else "MATCH"

        values = [
            shift_date.strftime("%d-%m-%Y"),  # col 1
            shift,                             # col 2
            shift_times.get(shift, ""),        # col 3
            db_b,                              # col 4
            db_p,                              # col 5
            cap_s,                             # col 6
            cap_m,                             # col 7
            sensor_p,                          # col 8
            sensor_b,                          # col 9
            dbl_img,                           # col 10
            neg_img,                           # col 11
            plate_diff,                        # col 12
            plate_status,                      # col 13
            bunch_diff,                        # col 14
            bunch_status,                      # col 15
        ]

        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.alignment = data_align
            cell.border = thin_border

            # Plate Diff coloring (col 12)
            if col == 12:
                if plate_diff > 0:
                    cell.fill = yellow_fill
                elif plate_diff < 0:
                    cell.fill = red_fill

            # Plate Status coloring (col 13)
            if col == 13:
                if plate_status == "DB MORE":
                    cell.fill = yellow_fill
                    cell.font = Font(bold=True, color='9C6500')
                elif plate_status == "CAPTURED MORE":
                    cell.fill = red_fill
                    cell.font = Font(bold=True, color='9C0006')
                elif plate_status == "MATCH":
                    cell.fill = green_fill
                    cell.font = Font(bold=True, color='006100')

            # Bunch Diff coloring (col 14)
            if col == 14:
                if bunch_diff > 0:
                    cell.fill = yellow_fill
                elif bunch_diff < 0:
                    cell.fill = red_fill

            # Bunch Status coloring (col 15)
            if col == 15:
                if bunch_status == "DB MORE":
                    cell.fill = yellow_fill
                    cell.font = Font(bold=True, color='9C6500')
                elif bunch_status == "CAPTURED MORE":
                    cell.fill = red_fill
                    cell.font = Font(bold=True, color='9C0006')
                elif bunch_status == "MATCH":
                    cell.fill = green_fill
                    cell.font = Font(bold=True, color='006100')

        row += 1

    # Summary / Totals row
    row += 1
    ws.cell(row=row, column=1, value="TOTALS").font = Font(bold=True)
    ws.cell(row=row, column=1).border = thin_border

    # Sum numeric columns (4-12, 14) — skip status text columns (13, 15)
    sum_cols = [4, 5, 6, 7, 8, 9, 10, 11, 12, 14]
    for col in range(2, 16):
        cell = ws.cell(row=row, column=col)
        if col in sum_cols:
            col_letter = get_column_letter(col)
            cell.value = f"=SUM({col_letter}5:{col_letter}{row - 2})"
            cell.font = Font(bold=True)
        cell.border = thin_border
        cell.alignment = data_align

    # Column widths (15 columns)
    col_widths = [14, 8, 16, 12, 16, 18, 18, 16, 16, 14, 12, 16, 14, 16, 14]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # -- Image Summary Sheet ------------------------------------------------
    ws2 = wb.create_sheet("Image Summary")
    ws2.merge_cells('A1:D1')
    ws2['A1'].value = "Image Folder Summary"
    ws2['A1'].font = Font(bold=True, size=14, color='2F5496')
    ws2['A1'].alignment = Alignment(horizontal='center')

    ws2.merge_cells('A2:D2')
    ws2['A2'].value = (
        f"Range: {from_dt.strftime('%d-%m-%Y %H:%M')} to "
        f"{to_dt.strftime('%d-%m-%Y %H:%M')} IST"
    )
    ws2['A2'].font = Font(italic=True, size=10)
    ws2['A2'].alignment = Alignment(horizontal='center')

    img_headers = ["Folder", "Count (in range)", "Total (all time)", "Notes"]
    for col, h in enumerate(img_headers, 1):
        cell = ws2.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for col_idx, vals in enumerate([
        ("doublePlateImages", dbl_in_range, dbl_total, "Plates detected as double"),
        ("NegImages (all subfolders)", neg_in_range, neg_total,
         "Negative/rejected images"),
    ], start=5):
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(row=col_idx, column=c, value=v)
            cell.border = thin_border
            if c in (2, 3):
                cell.alignment = data_align

    for col in range(1, 5):
        ws2.column_dimensions[get_column_letter(col)].width = 28

    # Save
    date_tag = f"{from_dt.strftime('%d%b')}_to_{to_dt.strftime('%d%b%Y')}"
    output_file = os.path.join(OUTPUT_DIR, f"report_{date_tag}.xlsx")
    wb.save(output_file)
    print(f"\n  Report saved to: {output_file}")
    print("=" * 60)


if __name__ == "__main__":
    generate_report()

    # # Date range only
    # generate_report("14-03-2026", "18-03-2026")

    # # Date + time range
    # generate_report("14-03-2026 06:00", "18-03-2026 22:00")

    # # With CSV fallback instead of DB
    # generate_report("14-03-2026", "18-03-2026", csv_path="../18mar_db.csv")